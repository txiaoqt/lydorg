"""
Authoritative Document Builder for Y-TRACE Black-Box Test Suites.
Formats DOCX files matching the official Y-TRACE test documentation standard.
"""

from dataclasses import dataclass, field
from pathlib import Path
from typing import List, Union

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor


@dataclass
class TestCase:
    id: str
    description: str
    preconditions: str
    steps: Union[List[str], str]
    test_data: str
    expected_result: str
    pass_fail_criteria: str
    actual_result: str = "Not Yet Tested"

    def formatted_steps(self) -> str:
        if isinstance(self.steps, list):
            return "\n".join(
                f"{i + 1}. {step}" if not step.strip().startswith(f"{i + 1}.") else step
                for i, step in enumerate(self.steps)
            )
        return self.steps.strip()


@dataclass
class TestGroup:
    number: Union[int, str]
    title: str
    test_cases: List[TestCase] = field(default_factory=list)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set cell margins in dxa (1 pt = 20 dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'  <w:top w:w="{top}" w:type="dxa"/>'
        f'  <w:left w:w="{left}" w:type="dxa"/>'
        f'  <w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'  <w:right w:w="{right}" w:type="dxa"/>'
        f"</w:tcMar>"
    )
    tcPr.append(tcMar)


def set_cell_shading(cell, color_hex: str):
    """Apply background shading color to a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    tcPr.append(shd)


def set_table_borders(table, color="B0BEC5", sz="4", val="single"):
    """Apply subtle, clean borders to a table."""
    tblPr = table._tbl.tblPr
    tblBorders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f"</w:tblBorders>"
    )
    tblPr.append(tblBorders)


def prevent_row_split(row):
    """Ensure a row does not split across a page break."""
    trPr = row._tr.get_or_add_trPr()
    cantSplit = parse_xml(f'<w:cantSplit {nsdecls("w")}/>')
    trPr.append(cantSplit)


def set_page_layout(doc: Document):
    section = doc.sections[0]
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.page_width = Inches(8.5)
    section.page_height = Inches(11.0)


def configure_styles(doc: Document):
    styles = doc.styles

    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor(17, 24, 39)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.line_spacing = 1.15


def add_title_page(doc: Document, section_title: str, section_number_str: str):
    # Top spacing
    for _ in range(3):
        doc.add_paragraph()

    # System header
    p_sys = doc.add_paragraph()
    p_sys.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    r_sys = p_sys.add_run("Y-TRACE System")
    r_sys.bold = True
    r_sys.font.name = "Calibri"
    r_sys.font.size = Pt(24)
    r_sys.font.color.rgb = RGBColor(30, 58, 138)  # Deep Navy

    # Module title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    r_title = p_title.add_run(section_title)
    r_title.bold = True
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(18)
    r_title.font.color.rgb = RGBColor(31, 41, 55)  # Dark Charcoal

    # Document Type
    p_type = doc.add_paragraph()
    p_type.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    r_type = p_type.add_run("Black-Box Test Cases")
    r_type.font.name = "Calibri"
    r_type.font.size = Pt(14)
    r_type.font.color.rgb = RGBColor(75, 85, 99)

    for _ in range(4):
        doc.add_paragraph()

    # Metadata Block
    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    meta_lines = [
        "System: Y-TRACE: Youth Organization Registration, Compliance, and Budget Monitoring System",
        f"Module Specification: Section {section_number_str} — {section_title}",
        "Testing Method: Black-Box Testing",
        "Basis: Current Y-TRACE Implementation (Codebase & UI Verified)",
        "Generated: September 16, 2026",
        "Updated: September 18, 2026",
        "Reviewed: September 18, 2026",
        "Status: Not Yet Tested",
    ]
    for i, line in enumerate(meta_lines):
        r = p_meta.add_run(line + ("\n" if i < len(meta_lines) - 1 else ""))
        r.font.name = "Calibri"
        r.font.size = Pt(10.5)
        if "Status:" in line:
            r.bold = True
            r.font.color.rgb = RGBColor(180, 83, 9)  # Warm Amber
        elif "System:" in line or "Module Specification:" in line:
            r.bold = True
            r.font.color.rgb = RGBColor(30, 58, 138)
        else:
            r.font.color.rgb = RGBColor(75, 85, 99)

    doc.add_page_break()


def add_table_of_contents(doc: Document, groups: List[TestGroup]):
    p_toc = doc.add_paragraph()
    r_toc = p_toc.add_run("Table of Contents")
    r_toc.bold = True
    r_toc.font.name = "Calibri"
    r_toc.font.size = Pt(16)
    r_toc.font.color.rgb = RGBColor(30, 58, 138)
    p_toc.paragraph_format.space_after = Pt(12)

    toc_table = doc.add_table(rows=1, cols=3)
    toc_table.style = "Table Grid"
    toc_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(toc_table, color="CBD5E1", sz="4")

    # Column widths
    col_widths = [Inches(0.9), Inches(4.3), Inches(1.3)]

    # Header row
    hdr_cells = toc_table.rows[0].cells
    hdr_titles = ["Section", "Features", "Test Cases"]
    for i, title in enumerate(hdr_titles):
        hdr_cells[i].text = title
        set_cell_shading(hdr_cells[i], "E8EEF5")
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=140, right=140)
        p = hdr_cells[i].paragraphs[0]
        if i != 1:
            p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        for run in p.runs:
            run.font.name = "Calibri"
            run.font.bold = True
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(30, 58, 138)

    prevent_row_split(toc_table.rows[0])

    total_cases = 0
    for group in groups:
        row = toc_table.add_row()
        prevent_row_split(row)
        cells = row.cells

        case_count = len(group.test_cases)
        total_cases += case_count
        if case_count > 0:
            first_id = group.test_cases[0].id
            last_id = group.test_cases[-1].id
            case_range = f"{first_id}-{last_id}" if case_count > 1 else first_id
        else:
            case_range = "N/A"

        cells[0].text = str(group.number)
        cells[1].text = group.title
        cells[2].text = case_range

        for i in range(3):
            set_cell_margins(cells[i], top=100, bottom=100, left=140, right=140)
            p = cells[i].paragraphs[0]
            if i != 1:
                p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
            for run in p.runs:
                run.font.name = "Calibri"
                run.font.size = Pt(10)

    # Total row
    tot_row = toc_table.add_row()
    prevent_row_split(tot_row)
    tot_cells = tot_row.cells
    tot_cells[0].text = ""
    tot_cells[1].text = "Total Test Cases"
    tot_cells[2].text = str(total_cases)
    set_cell_shading(tot_cells[1], "F1F5F9")
    set_cell_shading(tot_cells[2], "F1F5F9")
    for i in range(3):
        set_cell_margins(tot_cells[i], top=120, bottom=120, left=140, right=140)
        p = tot_cells[i].paragraphs[0]
        if i == 2:
            p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        for run in p.runs:
            run.font.name = "Calibri"
            run.font.bold = True
            run.font.size = Pt(10)

    # Set column widths across all rows
    for row in toc_table.rows:
        for i, w in enumerate(col_widths):
            row.cells[i].width = w

    p_summary = doc.add_paragraph()
    p_summary.paragraph_format.space_before = Pt(16)
    r_sum = p_summary.add_run(f"Total Test Cases: {total_cases}")
    r_sum.bold = True
    r_sum.font.name = "Calibri"
    r_sum.font.size = Pt(11)

    doc.add_page_break()


def render_test_case_table(doc: Document, tc: TestCase):
    table = doc.add_table(rows=8, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, color="CBD5E1", sz="4")

    col_widths = [Inches(1.8), Inches(4.7)]

    field_data = [
        ("Test Case ID", tc.id, True, "E8EEF5"),
        ("Description", tc.description, False, None),
        ("Preconditions", tc.preconditions, False, None),
        ("Test Steps", tc.formatted_steps(), False, None),
        ("Test Data", tc.test_data, False, None),
        ("Expected Result", tc.expected_result, False, None),
        ("Actual Result", tc.actual_result, False, "F1F5F9"),
        ("Pass/Fail Criteria", tc.pass_fail_criteria, False, None),
    ]

    for idx, (label, val, is_header, bg_color) in enumerate(field_data):
        row = table.rows[idx]
        prevent_row_split(row)
        cell_lbl, cell_val = row.cells[0], row.cells[1]

        cell_lbl.text = label
        cell_val.text = val

        set_cell_margins(cell_lbl, top=90, bottom=90, left=140, right=140)
        set_cell_margins(cell_val, top=90, bottom=90, left=140, right=140)

        # Apply header/row backgrounds
        if bg_color:
            set_cell_shading(cell_lbl, bg_color)
            set_cell_shading(cell_val, bg_color)

        # Style label paragraph
        p_lbl = cell_lbl.paragraphs[0]
        p_lbl.paragraph_format.space_before = Pt(0)
        p_lbl.paragraph_format.space_after = Pt(0)
        p_lbl.paragraph_format.line_spacing = 1.15
        for run in p_lbl.runs:
            run.font.name = "Calibri"
            run.font.bold = True
            run.font.size = Pt(9.5)
            if is_header:
                run.font.color.rgb = RGBColor(30, 58, 138)

        # Style value paragraph
        p_val = cell_val.paragraphs[0]
        p_val.paragraph_format.space_before = Pt(0)
        p_val.paragraph_format.space_after = Pt(0)
        p_val.paragraph_format.line_spacing = 1.15
        for run in p_val.runs:
            run.font.name = "Calibri"
            run.font.size = Pt(9.5)
            if is_header:
                run.font.bold = True
                run.font.color.rgb = RGBColor(30, 58, 138)
            elif label == "Actual Result":
                run.font.italic = True
                run.font.color.rgb = RGBColor(75, 85, 99)

        # Preserve line breaks formatting in steps and criteria
        if "\n" in val:
            cell_val.text = ""
            lines = val.split("\n")
            for line_idx, line in enumerate(lines):
                p = cell_val.add_paragraph() if line_idx > 0 else cell_val.paragraphs[0]
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(2 if line_idx < len(lines) - 1 else 0)
                p.paragraph_format.line_spacing = 1.15
                r = p.add_run(line)
                r.font.name = "Calibri"
                r.font.size = Pt(9.5)

    # Set column widths
    for row in table.rows:
        row.cells[0].width = col_widths[0]
        row.cells[1].width = col_widths[1]

    # Spacing between test cases
    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(0)
    p_spacer.paragraph_format.space_after = Pt(6)


def build_test_suite_document(
    section_number_str: str,
    section_title: str,
    groups: List[TestGroup],
    output_path: Union[str, Path],
):
    doc = Document()
    set_page_layout(doc)
    configure_styles(doc)

    # Title Page
    add_title_page(doc, section_title, section_number_str)

    # Table of Contents
    add_table_of_contents(doc, groups)

    # Test Groups & Cases
    for group in groups:
        p_grp = doc.add_paragraph()
        p_grp.paragraph_format.space_before = Pt(14)
        p_grp.paragraph_format.space_after = Pt(8)
        p_grp.paragraph_format.keep_with_next = True
        r_grp = p_grp.add_run(f"{group.number}. {group.title}")
        r_grp.bold = True
        r_grp.font.name = "Calibri"
        r_grp.font.size = Pt(13)
        r_grp.font.color.rgb = RGBColor(30, 58, 138)

        for tc in group.test_cases:
            render_test_case_table(doc, tc)

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    print(f"Successfully generated: {output_path.name} ({sum(len(g.test_cases) for g in groups)} test cases)")
