"""
Build, audit, verify, and package all 17 Y-TRACE Black-Box Test Case documents.
Generates to both docs/test-cases/ and docs/testcases/, then zips the updated package.
"""

import os
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path

# Add scripts directory to sys.path
scripts_dir = Path(__file__).parent.resolve()
workspace_root = scripts_dir.parent.resolve()
sys.path.insert(0, str(scripts_dir))

SECTION_MODULES = [
    ("01", "gen_section_01", "01_YTRACE_Authentication_Access_Black_Box_Test_Cases.docx"),
    ("02", "gen_section_02", "02_YTRACE_Admin_Authentication_RBAC_Black_Box_Test_Cases.docx"),
    ("03", "gen_section_03", "03_YTRACE_Terms_Privacy_Black_Box_Test_Cases.docx"),
    ("04", "gen_section_04", "04_YTRACE_Organization_Profile_YORP_Registry_Black_Box_Test_Cases.docx"),
    ("05", "gen_section_05", "05_YTRACE_Registration_Compliance_Black_Box_Test_Cases.docx"),
    ("06", "gen_section_06", "06_YTRACE_YORP_Renewal_Black_Box_Test_Cases.docx"),
    ("07", "gen_section_07", "07_YTRACE_YPOP_PPA_Black_Box_Test_Cases.docx"),
    ("08", "gen_section_08", "08_YTRACE_Budget_Request_Creation_Black_Box_Test_Cases.docx"),
    ("09", "gen_section_09", "09_YTRACE_Budget_Request_Review_Release_Black_Box_Test_Cases.docx"),
    ("10", "gen_section_10", "10_YTRACE_Liquidation_Reporting_Black_Box_Test_Cases.docx"),
    ("11", "gen_section_11", "11_YTRACE_Budget_Monitoring_Black_Box_Test_Cases.docx"),
    ("12", "gen_section_12", "12_YTRACE_Public_Budget_Transparency_Black_Box_Test_Cases.docx"),
    ("13", "gen_section_13", "13_YTRACE_Forms_Templates_Black_Box_Test_Cases.docx"),
    ("14", "gen_section_14", "14_YTRACE_News_Releases_Content_Black_Box_Test_Cases.docx"),
    ("15", "gen_section_15", "15_YTRACE_Inquiries_Notifications_Audit_Black_Box_Test_Cases.docx"),
    ("16", "gen_section_16", "16_YTRACE_Public_Information_Portal_Black_Box_Test_Cases.docx"),
    ("17", "gen_section_17", "17_YTRACE_Desktop_Cross_Surface_Regression_Black_Box_Test_Cases.docx"),
]


def generate_all_docs():
    print("=" * 80)
    print("STEP 1: GENERATING ALL 17 TEST SUITE DOCX FILES")
    print("=" * 80)

    dir_hyphen = workspace_root / "docs" / "test-cases"
    dir_nohyphen = workspace_root / "docs" / "testcases"

    dir_hyphen.mkdir(parents=True, exist_ok=True)
    dir_nohyphen.mkdir(parents=True, exist_ok=True)

    total_all_cases = 0
    generated_files = []

    for sec_num, mod_name, filename in SECTION_MODULES:
        mod = __import__(mod_name)
        # Call main or builder
        get_groups_func = getattr(mod, f"get_section_{sec_num}_groups")
        groups = get_groups_func()
        case_count = sum(len(g.test_cases) for g in groups)
        total_all_cases += case_count

        # Build directly to destination
        from test_doc_builder import build_test_suite_document
        # Determine title from module
        doc_titles = {
            "01": "Authentication, Registration & User Access Control Module",
            "02": "Administrative Authentication, Role-Based Access Control & Desktop Viewport Gating Module",
            "03": "Terms of Service & Privacy Policy Enforcement Module",
            "04": "Organization Profile & YORP Registry Module",
            "05": "Registration Compliance & Document Verification Module",
            "06": "YORP Accreditation Renewal Module",
            "07": "YPOP Incentive & PPA Submission Module",
            "08": "Budget Request Creation & Itemized Budgeting Module",
            "09": "Budget Request Review & Release Pipeline Module",
            "10": "Liquidation Reporting & Review Module",
            "11": "Budget Monitoring & Financial Tracking Module",
            "12": "Public Budget Transparency Portal Module",
            "13": "Forms & Templates Management (Archive / Restore) Module",
            "14": "News Releases & Public Content Module",
            "15": "Inquiries, Notifications, & Audit Logs Module",
            "16": "Public Information Portal Module",
            "17": "Desktop Cross-Surface & Integration Regression Module",
        }
        title = doc_titles.get(sec_num, f"Section {sec_num} Module")

        out_hyphen = dir_hyphen / filename
        out_nohyphen = dir_nohyphen / filename

        build_test_suite_document(
            section_number_str=sec_num,
            section_title=title,
            groups=groups,
            output_path=out_hyphen,
        )

        # Copy to nohyphen folder
        shutil.copy2(out_hyphen, out_nohyphen)

        generated_files.append((filename, case_count, len(groups)))
        print(f"  [+] Section {sec_num}: {filename} -> {case_count} test cases across {len(groups)} groups")

    # Remove any obsolete shortened DOCX files
    obsolete_files = [
        "01_YTRACE_Auth_Access_Black_Box_Test_Cases.docx",
        "02_YTRACE_Admin_Auth_RBAC_Black_Box_Test_Cases.docx",
        "04_YTRACE_Org_Profile_Registry_Black_Box_Test_Cases.docx",
    ]
    for obs in obsolete_files:
        for d in [dir_hyphen, dir_nohyphen]:
            obs_path = d / obs
            if obs_path.exists():
                obs_path.unlink()
                print(f"  [-] Removed obsolete shortened file: {obs_path}")

    print("-" * 80)
    print(f"Total Test Cases Generated Across All 17 Sections: {total_all_cases}")
    print("=" * 80)
    return generated_files, total_all_cases


def audit_and_verify_all():
    print("\n" + "=" * 80)
    print("STEP 2: RUNNING TABLE STRUCTURE AUDIT AND DOCUMENT VERIFICATION")
    print("=" * 80)

    from docx import Document

    dir_hyphen = workspace_root / "docs" / "test-cases"
    all_passed = True

    for sec_num, _, filename in SECTION_MODULES:
        filepath = dir_hyphen / filename
        if not filepath.exists():
            print(f"  [!] Missing file: {filepath}")
            all_passed = False
            continue

        doc = Document(str(filepath))
        # Table 0: Metadata table if any, or TOC table
        # Find TOC table
        toc_table = doc.tables[0]
        # Verify TOC rows
        toc_rows = len(toc_table.rows)

        # Test case tables: remaining tables
        tc_tables = doc.tables[1:]
        
        # Check every test case table has exactly 8 rows
        invalid_tables = []
        for idx, tbl in enumerate(tc_tables):
            if len(tbl.rows) != 8:
                invalid_tables.append((idx + 1, len(tbl.rows)))

        if invalid_tables:
            print(f"  [FAIL] {filename}: {len(invalid_tables)} tables do not have 8 rows: {invalid_tables[:5]}")
            all_passed = False
        else:
            print(f"  [PASS] {filename}: {len(tc_tables)} test case tables (all 8 rows exactly), TOC rows: {toc_rows}")

    return all_passed


def package_zip_archive():
    print("\n" + "=" * 80)
    print("STEP 3: PACKAGING ZIP ARCHIVE")
    print("=" * 80)

    zip_filename = "YTRACE_Black_Box_Test_Cases_Updated_2026-09-18_Final.zip"
    
    dir_hyphen = workspace_root / "docs" / "test-cases"
    dir_nohyphen = workspace_root / "docs" / "testcases"

    zip_path_hyphen = dir_hyphen / zip_filename
    zip_path_nohyphen = dir_nohyphen / zip_filename

    # Create zip from hyphen directory
    with zipfile.ZipFile(zip_path_hyphen, "w", zipfile.ZIP_DEFLATED) as zf:
        for sec_num, _, filename in SECTION_MODULES:
            src_file = dir_hyphen / filename
            zf.write(src_file, arcname=filename)
            print(f"  [+] Added to zip: {filename} ({src_file.stat().st_size:,} bytes)")

    shutil.copy2(zip_path_hyphen, zip_path_nohyphen)
    print(f"  [+] Packaged archive: {zip_path_hyphen} ({zip_path_hyphen.stat().st_size:,} bytes)")
    print(f"  [+] Mirrored archive: {zip_path_nohyphen}")
    print("=" * 80)


def main():
    generated_files, total_cases = generate_all_docs()
    audit_passed = audit_and_verify_all()
    if not audit_passed:
        print("\n[!] Audit reported failures. Please check above logs.")
        sys.exit(1)
    
    package_zip_archive()
    print("\nSUCCESS: All 17 Black-Box manual test documents generated, verified, and packaged successfully.")


if __name__ == "__main__":
    main()
