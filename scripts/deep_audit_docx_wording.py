import os
import glob
import re
from docx import Document

KEYWORDS = [
    r'\bDOM\b',
    r'\bRPC\b',
    r'\bRPCs\b',
    r'\bRLS\b',
    r'\bforeign key\b',
    r'\bforeign keys\b',
    r'\bdatabase row\b',
    r'\bdatabase rows\b',
    r'\bdatabase table\b',
    r'\bdatabase tables\b',
    r'\bdatabase transaction\b',
    r'\bdatabase records\b',
    r'\bschema\b',
    r'\bschemas\b',
    r'\bsession state\b',
    r'\bsession token\b',
    r'\bsession tokens\b',
    r'\btoken-authenticated\b',
    r'\bToken-Authenticated\b',
    r'\brouting guard\b',
    r'\brouting guards\b',
    r'\bauthentication guard\b',
    r'\bauthentication guards\b',
    r'\brouter guard\b',
    r'\brouter guards\b',
    r'\bquery string\b',
    r'\bquery strings\b',
    r'\bforged\b',
    r'REG-YYYY-NNNN',
    r'internal registration reference',
    r'bypass authentication',
    r'Observe application routing behavior',
    r'inspect exported schema',
    r'canonical table rendering and schema',
]

print("=== SCANNING DOCX FILES ===")
for docx_path in sorted(glob.glob('docs/testcases/*.docx')):
    doc = Document(docx_path)
    found = []
    
    # Check paragraphs
    for p_idx, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        for kw in KEYWORDS:
            if re.search(kw, text, re.IGNORECASE):
                found.append((f"P{p_idx}", kw, text))
                
    # Check tables
    for t_idx, table in enumerate(doc.tables):
        for r_idx, row in enumerate(table.rows):
            for c_idx, cell in enumerate(row.cells):
                cell_text = cell.text.strip()
                for kw in KEYWORDS:
                    if re.search(kw, cell_text, re.IGNORECASE):
                        found.append((f"T{t_idx} R{r_idx} C{c_idx}", kw, cell_text))
                        
    # Deduplicate matches
    unique_found = []
    seen = set()
    for loc, kw, text in found:
        key = (kw, text[:60])
        if key not in seen:
            seen.add(key)
            unique_found.append((loc, kw, text))
            
    if unique_found:
        print(f"\n[{os.path.basename(docx_path)}] ({len(unique_found)} unique matches):")
        for loc, kw, text in unique_found:
            clean_text = text.encode('ascii', 'replace').decode('ascii').replace('\n', ' ')
            print(f"  {loc:15s} [{kw}]: {clean_text[:120]}")
