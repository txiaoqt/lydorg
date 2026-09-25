"""
Master Verification Script for the 17 Y-TRACE Black-Box Test Suite Documents.
Accurately validates TOC summary table (Table 0) and Test Case tables (Tables 1..N).
"""

import sys
from pathlib import Path
from docx import Document

workspace_root = Path(__file__).parent.parent.resolve()
target_dir = workspace_root / "docs" / "test-cases"

SECTIONS = [
    ("01", "01_YTRACE_Authentication_Access_Black_Box_Test_Cases.docx", "Authentication & Access Control"),
    ("02", "02_YTRACE_Admin_Authentication_RBAC_Black_Box_Test_Cases.docx", "Administrative Authentication, Role-Based Access Control & Desktop Viewport Gating"),
    ("03", "03_YTRACE_Terms_Privacy_Black_Box_Test_Cases.docx", "Terms of Service & Privacy Policy Enforcement"),
    ("04", "04_YTRACE_Organization_Profile_YORP_Registry_Black_Box_Test_Cases.docx", "Organization Profile & YORP Registry"),
    ("05", "05_YTRACE_Registration_Compliance_Black_Box_Test_Cases.docx", "Registration Document Submission & Review (Compliance)"),
    ("06", "06_YTRACE_YORP_Renewal_Black_Box_Test_Cases.docx", "YORP Accreditation Renewal Workflow"),
    ("07", "07_YTRACE_YPOP_PPA_Black_Box_Test_Cases.docx", "YPOP Incentive & PPA Submission"),
    ("08", "08_YTRACE_Budget_Request_Creation_Black_Box_Test_Cases.docx", "Budget Request Creation & Management"),
    ("09", "09_YTRACE_Budget_Request_Review_Release_Black_Box_Test_Cases.docx", "Budget Request Review & Release Pipeline"),
    ("10", "10_YTRACE_Liquidation_Reporting_Black_Box_Test_Cases.docx", "Liquidation Reporting & Review"),
    ("11", "11_YTRACE_Budget_Monitoring_Black_Box_Test_Cases.docx", "Budget Monitoring & Financial Tracking"),
    ("12", "12_YTRACE_Public_Budget_Transparency_Black_Box_Test_Cases.docx", "Public Budget Transparency Portal"),
    ("13", "13_YTRACE_Forms_Templates_Black_Box_Test_Cases.docx", "Forms & Templates Management (Archive / Restore)"),
    ("14", "14_YTRACE_News_Releases_Content_Black_Box_Test_Cases.docx", "News Releases & Public Content"),
    ("15", "15_YTRACE_Inquiries_Notifications_Audit_Black_Box_Test_Cases.docx", "Inquiries, Notifications, & Audit Logs"),
    ("16", "16_YTRACE_Public_Information_Portal_Black_Box_Test_Cases.docx", "Public Information Portal"),
    ("17", "17_YTRACE_Desktop_Cross_Surface_Regression_Black_Box_Test_Cases.docx", "Desktop Cross-Surface & Integration Regression"),
]


def audit_document(file_path: Path):
    if not file_path.exists():
        return {
            "exists": False,
            "error": "File does not exist",
            "tc_count": 0,
            "errors": ["File does not exist"],
        }

    doc = Document(file_path)
    errors = []

    if len(doc.tables) < 2:
        return {
            "exists": True,
            "error": "Document does not have sufficient tables (needs TOC + test cases)",
            "tc_count": 0,
            "errors": ["Insufficient tables"],
        }

    # Table 0 is Table of Contents
    toc_table = doc.tables[0]
    toc_headers = [c.text.strip() for c in toc_table.rows[0].cells]
    if toc_headers != ["Section", "Features", "Test Cases"]:
        errors.append(f"TOC header mismatch: {toc_headers}")

    # Total row in TOC
    last_row = toc_table.rows[-1]
    toc_total_str = last_row.cells[2].text.strip()
    try:
        toc_total = int(toc_total_str)
    except ValueError:
        toc_total = -1
        errors.append(f"Invalid TOC total: '{toc_total_str}'")

    # Tables 1..N are Test Cases
    tc_tables = doc.tables[1:]
    actual_tc_count = len(tc_tables)

    if toc_total != -1 and toc_total != actual_tc_count:
        errors.append(f"TOC total ({toc_total}) does not match actual test case tables count ({actual_tc_count})")

    # Expected field labels in rows 0..7
    expected_labels = [
        "Test Case ID",
        "Description",
        "Preconditions",
        "Test Steps",
        "Test Data",
        "Expected Result",
        "Actual Result",
        "Pass/Fail Criteria",
    ]

    # Audit each test case table
    for idx, table in enumerate(tc_tables, 1):
        if len(table.rows) != 8:
            errors.append(f"TC Table {idx} has {len(table.rows)} rows (expected 8)")
            continue

        for r_idx, label in enumerate(expected_labels):
            cell_lbl = table.rows[r_idx].cells[0].text.strip()
            if cell_lbl != label:
                errors.append(f"TC Table {idx} Row {r_idx} label is '{cell_lbl}' (expected '{label}')")

        # Check Test Case ID in Row 0 Cell 1
        tc_id = table.rows[0].cells[1].text.strip()
        if not tc_id.startswith("TC"):
            errors.append(f"TC Table {idx} ID '{tc_id}' does not start with 'TC'")

        # Check Actual Result in Row 6 Cell 1
        actual_res = table.rows[6].cells[1].text.strip()
        if actual_res != "Not Yet Tested":
            errors.append(f"TC Table {idx} Actual Result is '{actual_res}' (expected 'Not Yet Tested')")

    return {
        "exists": True,
        "tc_count": actual_tc_count,
        "toc_total": toc_total,
        "errors": errors,
    }


def main():
    print("=" * 85)
    print("MASTER QA AUDIT REPORT — Y-TRACE 17-SECTION BLACK-BOX TEST SUITE")
    print("=" * 85)

    total_suite_cases = 0
    all_passed = True

    for sec_num, filename, sec_name in SECTIONS:
        file_path = target_dir / filename
        res = audit_document(file_path)

        if not res["exists"] or res["errors"]:
            all_passed = False
            status_str = f"FAILED ({len(res['errors'])} errors)"
            print(f"[{sec_num}] {filename} -> {status_str}")
            for err in res["errors"][:5]:
                print(f"     * {err}")
        else:
            total_suite_cases += res["tc_count"]
            print(f"[{sec_num}] {filename} -> PASSED ({res['tc_count']} test cases)")

    print("=" * 85)
    print(f"TOTAL TEST CASES ACROSS ALL 17 SECTIONS: {total_suite_cases}")
    print(f"AUDIT INTEGRITY STATUS: {'ALL 17 DOCUMENTS PASSED' if all_passed else 'ERRORS FOUND'}")
    print("=" * 85)


if __name__ == "__main__":
    main()
